from pathlib import Path
import json, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path('Audit_2026-10-26/Independent_Readiness_Audit')
audit=json.loads((root/'audit-findings.json').read_text())
coverage=json.loads((root/'page-coverage.json').read_text())
sources=[('Vanta: audit workflows and collaboration','https://www.vanta.com/products/audit'),('Vanta: audit evidence states','https://help.vanta.com/en/articles/11345427-audit-evidence'),('Vanta: policy lifecycle','https://help.vanta.com/en/articles/11345363-getting-started-with-policies'),('Vanta: risk assessments','https://help.vanta.com/en/articles/15089313-completing-risk-assessments'),('Vanta: product feature reference','https://www.vanta.com/features'),('NIST: conducting risk assessments','https://csrc.nist.gov/pubs/sp/800/30/r1/final'),('ISO: ISO/IEC 27001 overview','https://committee.iso.org/standard/27001')]
benchmarks=[
('Overview and next actions','The dashboard gives labelled status cards and assessment next steps. Keep recorded status distinct from verified evidence.','Useful foundation; applicability accuracy needs H3.'),
('Policies and employee acceptance','Vanta links drafting, approval, published policies and employee acceptance. ComplianceCore already has corresponding versioned approval/readership records.','Broadly aligned in concept; retain the stricter independent-review policy where chosen.'),
('Risk ownership and assessment','Vanta supports owners preparing assessments and designated approvers reviewing them. ComplianceCore has asset scenarios, ratings, treatment links and independent reviews, but preparation is coordinator-led.','Add explicit risk-owner participation/acceptance and organisation-approved scoring/acceptance criteria. Do not equate a free-text note with every required governance decision.'),
('Audit workspace and evidence review','Vanta provides audit-scoped evidence, requests, review states and collaboration. ComplianceCore stores proof and review decisions, but its fresh audit journey is blocked and collaboration is more limited.','Fix H1 first. Add a coherent audit workspace with request, submission, return/comment and approval context, with deliberate auditor sharing scope.'),
('People and training','ComplianceCore has programme and verification concepts but lacks a connected participant-assignment journey.','Fix H2 and M1 before using it as the training compliance record.'),
('Evidence and automation','Vanta documents connected evidence collection and continuing tests. ComplianceCore has manual uploads, version-pinned proof, APIs and webhooks; broad automated collection was not established in this audit.','Describe manual and automated coverage accurately. Add only verified connectors; show last collection time, collection failures and evidence age.'),
('Roles and access','ComplianceCore has organisation-bound roles and API checks. Some UI actions mismatch them, and ownership changes have a concurrency defect in the service probe.','Close H4 and M3. Use real multi-user acceptance tests and granular audit sharing before enterprise rollout.'),
('Framework-specific completeness','A framework catalogue and a generic workflow do not establish complete/current coverage for privacy, financial, continuity or security obligations. Existing profile notes warn about older catalogue editions.','Validate licensed/current requirements and organisation applicability. Retain scheme-specific evidence and do not promise certification or automatic legal compliance.')]
module_notes={
'dashboard':'Clearer layout; current progress remains dependent on approved applicability (H3).',
'frameworks':'Scope -> duties -> risks -> treatment -> applicability ordering is sensible. Verify full requirements coverage and specialist review.',
'risks':'Scenario, asset, likelihood, impact and treatment inputs are useful. Add explicit owner governance and approved scoring criteria.',
'policies':'Draft, approval, publication and readership concepts are present; labels need M4 review.',
'policy-templates':'Useful starting documents; tailor placeholders and approve before use. Downloads are not proof of implementation.',
'controls':'H3: ordinary N/A updates bypass the intended applicability decision path.',
'evidence':'Versioned proof is useful. Evidence rechecked after a harness selector error; fix nested main and label semantics (M4).',
'audits':'H1: summary input and the completion action do not support the service requirements.',
'vendors':'Questionnaire asks relevant service/data/assurance/contract/continuity questions. Improve explicit accountable-owner assignment and field labels.',
'incidents':'Decision-specific notification facts and verified deadlines are better than assuming a universal deadline. A qualified person must confirm the duty.',
'training':'H2 and M1: participant assignment is missing; blank optional frequency fails client validation.',
'calendar':'M2: end/start ordering validation is missing.',
'compliance-score':'M3: Take Snapshot is offered to roles without the required write permission. Keep scoring methodology explicit.',
'settings/notifications':'M3: save is offered to a read-only compliance manager.',
'settings/team':'H4: serialise ownership mutations. Existing role descriptions and ordinary API guards are useful.',
'integrations':'API/webhooks are not equivalent to a verified continuous evidence-collection platform. Keep coming-soon capabilities explicit.',
'reports':'Useful management view; reconcile metric definitions and review underlying proof before relying on a conclusion.',
'analytics':'Charts summarise records and statuses; they do not independently verify operation.'}
limits=[
'This is a product/workflow and engineering readiness review, using publicly available Vanta documentation as a benchmark. It is not an independent certification, legal opinion, penetration test or a claim of Vanta equivalence.',
'Current-source inspection covered 134 TSX files, including 484 field/label nodes and 90 JSX link nodes. Static enumeration does not validate every dynamic destination or every possible field value.',
'All 37 principal routes rendered in a local current-source preview against the managed demo API, with writes blocked, plus nine workflow panels. Seven create/upload forms were inspected across that run and the evidence recheck. A selector issue caused the first evidence snapshot to fail; its dedicated recheck rendered successfully.',
'Fresh browser checks of the canonical online frontend timed out before reaching sign-in fields. The frontend project still pointed to the expected deployment; HTML and the entry script were retrieved, but delivery was intermittently slow. This audit does not attribute those timeouts conclusively to an application defect or declare the live site freshly passed.',
'The earlier 86 live browser checks, 42 HTTP checks, download checks and reload checks remain historical functional evidence for this source release. They do not prove that a newly created organisation can complete every workflow. Demo records prepared by a seed script can hide UI gaps.',
'Six isolated probes reproduced current deficiencies: a control-owner N/A change, the missing audit summary field, full-payload completion rejection, reversed audit dates, reversed calendar dates and concurrent co-owner demotion. Passing these probes means the defect was observed; it is not a readiness pass.',
'The ownership race used controlled repository fixtures, not simultaneous writes to a production database. Date probes checked schemas/services rather than a live insert. The training optional-number issue was reproduced by executing the actual client schema locally.',
'No application change, deployment, production business-record mutation, external message or real payment was made. Sign-in/sign-out created normal authentication-session activity. New-user onboarding, real payment/email/provider delivery, accessibility with assistive hardware, load/concurrency at scale, and full object-store disaster recovery require separate acceptance evidence.'
]
acceptance=[
'Close H1-H4 with targeted regression tests and a real isolated-database concurrency test.',
'Close M1-M4: optional-field handling, chronological validation, role-correct actions and accessible fields/dialogs.',
'Run a fresh-organisation lifecycle through scope, duties, risk ownership/treatment, applicability, approved policies/readership, uploads and version changes, passing/failing tests, corrective action, audit report review/completion, training enrolment and management reporting.',
'Test each role with separate real identities, including negative API calls, ownership changes, session revocation, record scope and independent-review restrictions. UI role simulation alone is insufficient.',
'Validate framework editions and full requirement mapping for the exact scope with a competent compliance lead. Record risk acceptance criteria and authorised risk-owner decisions.',
'Confirm protected uploads, retention/access configuration, recovery of both database and stored evidence, operational monitoring, provider delivery and the intended business support arrangements.',
'Retest the canonical online deployment on desktop and phone under stable connectivity, including actual downloads and saved-record lifecycles. Resolve or explain delivery timeouts before customer rollout.'
]
doc=Document(); sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10);style.paragraph_format.space_after=Pt(6)
for n in ['Title','Heading 1','Heading 2']:
 doc.styles[n].font.name='Calibri';doc.styles[n].font.color.rgb=RGBColor.from_string('086963')
doc.add_heading('ComplianceCore\nCompliance Readiness Audit',0)
doc.add_paragraph('26 October 2026 | Vanta-referenced product and workflow review')
p=doc.add_paragraph();run=p.add_run(audit['verdict']);run.bold=True;run.font.color.rgb=RGBColor.from_string('B42318')
doc.add_paragraph('The application is a useful foundation and a demonstrable prototype. I would not approve unrestricted production use or represent it as a complete enterprise compliance system yet. Four high-priority workflow/access issues and four medium-priority input/usability issues remain. Use synthetic data for demonstrations until the blocking issues are closed; a real-data pilot needs an explicit, bounded risk decision and verified controls.')
doc.add_paragraph('Live address: '+audit['liveUrl']);doc.add_paragraph('Reviewed runtime source: '+audit['reviewedSource'])
doc.add_heading('What is working conceptually',1)
for s in ['The simplified dashboard, grouped navigation and role guidance are clearer than the earlier arrangement.','Framework assessment follows scope, duties, risk assessment, treatment and later applicability. Risks remain distinct from requirements coverage.','Policies have approval/publication and exact-version readership concepts. Evidence and several review workflows pin supporting versions.','The product distinguishes internal assessment review from external certification or attestation. Implementation of safeguards remains a legitimate part of compliance work.']:
 doc.add_paragraph(s,style='List Bullet')
doc.add_heading('Confirmed findings and acceptance checks',1)
for f in audit['findings']:
 doc.add_heading(f"{f['id']} | {f['severity']} | {f['title']}",2)
 doc.add_paragraph(f['observation']);doc.add_paragraph('Required correction: '+f['fix']);doc.add_paragraph('Closure evidence: '+f['acceptance'])
 doc.add_paragraph('Source evidence: '+'; '.join(f"{e['file']}:{e['line']}" for e in f['evidence']))
doc.add_heading('Vanta benchmark and product priorities',1)
doc.add_paragraph('Vanta is a reference for coherent workflows and collaboration, not a requirement to copy every feature, approval rule or paid-plan capability. The comparison below uses public documentation only; no private Vanta workspace was inspected.')
for title,comparison,action in benchmarks:
 doc.add_heading(title,2);doc.add_paragraph(comparison);doc.add_paragraph('Priority: '+action)
doc.add_heading('Page-by-page review coverage',1)
doc.add_paragraph('Status below means the current local UI rendered with the managed demo API. It does not mean the complete saved-data lifecycle passed. Fields and linked actions were also reviewed in source. See the evidence inventory for exact observations.')
table=doc.add_table(rows=1, cols=3);table.style='Light Shading Accent 1'
for c,t in zip(table.rows[0].cells,['Page','Inspection','Assessment']):c.text=t
for row in coverage['pages']:
 route=row['route'];cells=table.add_row().cells;cells[0].text='/'+route;cells[1].text='Rendered (read-only)' if row.get('opened') else 'Not verified';cells[2].text=module_notes.get(route,'No additional blocking defect confirmed in this limited inspection; this does not certify every input, role, or saved-data lifecycle.')
for row in table.rows:
 trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
doc.add_heading('Conditions before routine use',1)
for a in acceptance:doc.add_paragraph(a,style='List Number')
doc.add_heading('Scope, limitations and evidence',1)
for l in limits:doc.add_paragraph(l,style='List Bullet')
doc.add_paragraph('Evidence directory: '+str(root.resolve()))
doc.add_paragraph('Key files: audit-findings.json; source-input-and-link-inventory.json; local-page-and-input-inventory.json; evidence-recheck.json; page-coverage.json; service-probe-results.json; training-input-probe.json. The copied probe file documents deficiencies and should not be treated as an application acceptance suite. It can be temporarily copied into backend/test/unit to rerun in the existing test harness.')
doc.add_heading('Public reference sources',1)
for title,url in sources:doc.add_paragraph(title+'\n'+url)
footer=sec.footer.paragraphs[0];footer.text='ComplianceCore | Internal readiness review | 26 October 2026';footer.style='Caption'
name='ComplianceCore_Compliance_Readiness_Audit_2026-10-26.docx';doc.save(root/name)
text=['COMPLIANCECORE READINESS AUDIT - 26 OCTOBER 2026',audit['verdict'],'','Four High and four Medium confirmed findings. No product fixes or deployment made during this audit.','']
for f in audit['findings']:text.extend([f"{f['id']} [{f['severity']}] {f['title']}",f['observation'],'Correction: '+f['fix'],'Acceptance: '+f['acceptance'],''])
text.extend(['VANTA BENCHMARK',''])
for title,comparison,action in benchmarks:text.extend([title,comparison,action,''])
text.extend(['SCOPE AND LIMITATIONS','']+limits+['','BEFORE USE']+acceptance+['','PUBLIC SOURCES'])
for title,url in sources:text.extend([title,url])
(root/'Readiness_Audit_Summary.txt').write_text('\n'.join(text),encoding='utf-8')
downloads=Path.home()/'Downloads';shutil.copy2(root/name,downloads/name);shutil.copy2(root/'Readiness_Audit_Summary.txt',downloads/'ComplianceCore_Readiness_Audit_2026-10-26.txt')
print(json.dumps({'wordReport':str(downloads/name),'findings':len(audit['findings']),'high':4,'medium':4,'pages':len(coverage['pages'])}))
